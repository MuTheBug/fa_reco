"""Excel (.xlsx) export: right-to-left sheet, typed cells, A4 print setup."""

from __future__ import annotations

from datetime import date, datetime, time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import Report

FONT = "Arial"
FILLS = {"bad": "FDE2E2", "warn": "FFF1D6", "info": "E3EEFC", "ok": "E3F5E8", "muted": "F1F3F5"}
THIN = Side(style="thin", color="9AA5B1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUMBER_FORMATS = {
    "date": "yyyy/mm/dd",
    "time": "hh:mm",
    "datetime": "yyyy/mm/dd hh:mm",
    "minutes": "[h]:mm",
    "percent": "0%",
    "float": "0.00",
    "hours": "0.0",
}


def _cell_value(value, kind):
    if value is None or value == "":
        return None
    if kind == "minutes":
        return value / 1440 if value else None
    if kind == "time" and isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if kind == "time" and isinstance(value, str) and ":" in value:
        h, m = value.split(":")[:2]
        return time(int(h), int(m))
    if kind == "date" and isinstance(value, datetime):
        return value.date()
    if kind == "datetime" and isinstance(value, datetime):
        return value.replace(microsecond=0)
    if isinstance(value, (int, float, date, datetime, time)):
        return value
    return str(value)


def _safe_title(title: str) -> str:
    bad = '[]:*?/\\'
    t = "".join(ch for ch in title if ch not in bad).strip() or "تقرير"
    return t[:31]


def export_xlsx(report: Report, path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = _safe_title(report.title)
    ws.sheet_view.rightToLeft = True
    cols = report.columns
    ncol = max(1, len(cols))
    last = get_column_letter(ncol)

    def banner(row, text, size, bold=True, color="0F2C4C"):
        ws.merge_cells(f"A{row}:{last}{row}")
        c = ws[f"A{row}"]
        c.value = text
        c.font = Font(name=FONT, size=size, bold=bold, color=color)
        c.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)

    row = 1
    meta = report.meta
    if meta.get("org"):
        banner(row, meta["org"], 14)
        row += 1
    if meta.get("org_subtitle"):
        banner(row, meta["org_subtitle"], 11, bold=False, color="374151")
        row += 1
    banner(row, report.title, 13)
    row += 1
    if report.subtitle:
        banner(row, report.subtitle, 11, bold=False, color="1F2937")
        row += 1
    row += 1

    header_row = row
    head_fill = PatternFill("solid", fgColor="DDE6F0")
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i, value=c.title.replace("\n", " "))
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = head_fill
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True,
                                   readingOrder=2)
    ws.row_dimensions[row].height = 30
    row += 1
    first_data = row
    widths = [max(8, min(40, len(c.title) * 1.3)) for c in cols]

    def write_row(values: dict, style: str = ""):
        nonlocal row
        tones = values.get("_tones") or {}
        for i, c in enumerate(cols, start=1):
            v = _cell_value(values.get(c.key), c.kind)
            cell = ws.cell(row=row, column=i, value=v)
            cell.border = BORDER
            cell.font = Font(name=FONT, size=10, bold=style in ("subtotal", "total"),
                             color="6B7280" if values.get("_tone") == "muted" else "111827")
            if c.kind in NUMBER_FORMATS:
                cell.number_format = NUMBER_FORMATS[c.kind]
            numeric = c.numeric or c.kind in ("time", "date", "datetime")
            cell.alignment = Alignment(horizontal="center" if numeric else "right",
                                       vertical="center", wrap_text=not numeric, readingOrder=2)
            tone = tones.get(c.key) or values.get("_tone")
            if style == "total":
                cell.fill = head_fill
            elif style == "subtotal":
                cell.fill = PatternFill("solid", fgColor="F2F4F7")
            elif tone in FILLS:
                cell.fill = PatternFill("solid", fgColor=FILLS[tone])
            if v is not None:
                text = v.strftime("%Y/%m/%d %H:%M") if isinstance(v, datetime) else str(v)
                widths[i - 1] = max(widths[i - 1], min(45, len(text) * 1.15 + 2))
        row += 1

    def labelled(values: dict, label: str) -> dict:
        out = dict(values)
        for c in cols:
            if not c.numeric and out.get(c.key) in (None, ""):
                out[c.key] = label
                break
        return out

    if report.sections:
        sec_fill = PatternFill("solid", fgColor="E9EEF5")
        for sec in report.sections:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
            cell = ws.cell(row=row, column=1, value=sec.title)
            cell.font = Font(name=FONT, bold=True, size=10, color="0F2C4C")
            cell.fill = sec_fill
            cell.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
            row += 1
            for v in sec.rows:
                write_row(v)
            if sec.subtotal and any(x not in (None, "", 0) for x in sec.subtotal.values()):
                write_row(labelled(sec.subtotal, "المجموع"), "subtotal")
    else:
        for v in report.rows:
            write_row(v)
    if report.totals and any(x not in (None, "", 0) for x in report.totals.values()):
        write_row(labelled(report.totals, "المجموع الكلي"), "total")

    if report.notes:
        row += 1
        for n in report.notes:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
            c = ws.cell(row=row, column=1, value=n)
            c.font = Font(name=FONT, size=9, color="374151")
            c.alignment = Alignment(horizontal="right", wrap_text=True, readingOrder=2)
            row += 1

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=first_data, column=1)
    if not report.sections and report.rows:
        ws.auto_filter.ref = f"A{header_row}:{last}{first_data + len(report.rows) - 1}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if report.landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "صفحة &P من &N"
    ws.oddHeader.center.text = report.title
    wb.properties.title = report.title
    wb.properties.creator = "Hadir"
    wb.save(path)
