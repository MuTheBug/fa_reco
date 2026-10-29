"""Employees: list, add/edit, face enrollment, Excel import."""

from __future__ import annotations

import cv2
import numpy as np
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QSize, Qt, QTimer
from PySide6.QtGui import QImage, QPixmap
from PySide6.QtWidgets import (QCheckBox, QComboBox, QDialog, QFileDialog, QFormLayout, QGridLayout,
                               QHBoxLayout, QLabel, QLineEdit, QListWidget, QListWidgetItem,
                               QProgressBar, QTabWidget, QVBoxLayout, QWidget)

from ...core.settings import SECURITY_LEVELS
from ...core.timeutil import digits, parse_date
from ...db.store import StoreError
from ...face.enroll import STEPS, chip_from_png
from ...face.service import duplicate_candidates, save_captures
from ..admin import Page
from ..theme import C, icon
from ..widgets.camera_view import CameraView
from ..widgets.common import (Card, DataTable, DateEdit, PageHeader, SearchBox, button, confirm,
                              error, info, label)

POSE_LABELS = {"front": "أمامي", "left": "يسار", "right": "يمين", "up": "للأعلى", "down": "للأسفل",
               "photo": "صورة", "auto": "تحديث تلقائي"}
IMPORT_HEADERS = ["الاسم", "رقم الهاتف", "الصفة الوظيفية", "المجموعة"]
# Column names used by the import template of version 1.0.
HEADER_ALIASES = {"الهاتف": "رقم الهاتف", "العنوان الوظيفي": "الصفة الوظيفية"}


def cell_text(v) -> str:
    """Excel cell -> text (whole numbers lose the trailing «.0» Excel gives them)."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return "" if v is None else str(v).strip()


class EmployeesPage(Page):
    def __init__(self, ctx):
        super().__init__(ctx)
        header = PageHeader("الموظفون", "إضافة الموظفين وتسجيل بصمة الوجه لكل منهم")
        self.lay.addWidget(header)
        add = button("إضافة موظف", "user_plus", "primary")
        add.clicked.connect(lambda: self.edit(None))
        header.add(add)
        imp = button("استيراد من Excel", "upload")
        imp.clicked.connect(self.import_excel)
        header.add(imp)
        tpl = button("نموذج الاستيراد", "download", "ghost")
        tpl.clicked.connect(self.save_template)
        header.add(tpl)

        card = Card()
        tools = QHBoxLayout()
        self.search = SearchBox("بحث بالاسم أو الرقم أو الصفة الوظيفية…")
        self.group = QComboBox()
        self.state = QComboBox()
        self.state.addItem("الفعّالون", True)
        self.state.addItem("غير الفعّالين", False)
        self.state.addItem("الكل", None)
        tools.addWidget(self.search, 1)
        tools.addWidget(label("المجموعة:"))
        tools.addWidget(self.group)
        tools.addWidget(self.state)
        card.lay.addLayout(tools)
        self.table = DataTable([
            ("code", "الرقم الوظيفي", "text"), ("full_name", "الاسم", "text"),
            ("job_title", "الصفة الوظيفية", "text"), ("phone", "رقم الهاتف", "text"),
            ("group_name", "المجموعة", "text"), ("face", "بصمة الوجه", "text"),
            ("state", "الحالة", "text")], stretch="full_name")
        self.table.setMinimumHeight(420)
        self.table.activated_row.connect(lambda r: self.edit(r["id"]))
        card.lay.addWidget(self.table)
        actions = QHBoxLayout()
        for text, ic, fn, kind in (("تعديل", "edit", self._edit_selected, ""),
                                   ("تسجيل الوجه", "face", self._enroll_selected, ""),
                                   ("السجل الزمني", "clock", self._timeline_selected, ""),
                                   ("حذف", "trash", self._delete_selected, "danger")):
            b = button(text, ic, kind)
            b.clicked.connect(fn)
            actions.addWidget(b)
        actions.addStretch(1)
        self.count = label("", "Hint")
        actions.addWidget(self.count)
        card.lay.addLayout(actions)
        self.lay.addWidget(card, 1)
        self.search.textChanged.connect(self.reload)
        self.group.currentIndexChanged.connect(self.reload)
        self.state.currentIndexChanged.connect(self.reload)

    def refresh(self):
        current = self.group.currentData()
        self.group.blockSignals(True)
        self.group.clear()
        self.group.addItem("الكل", None)
        for g in self.db.groups.list(include_inactive=True):
            self.group.addItem(g["name"], g["id"])
        self.group.setCurrentIndex(max(0, self.group.findData(current)))
        self.group.blockSignals(False)
        self.reload()

    def reload(self):
        rows = []
        for e in self.db.employees.list(self.search.text().strip(), self.group.currentData(),
                                         self.state.currentData()):
            r = dict(e)
            n = e["samples"]
            r["face"] = f"{digits(n)} عينة" if n else "غير مسجّل"
            r["state"] = "فعّال" if e["active"] else "غير فعّال"
            r["_tones"] = {"face": "warn" if not n else ("ok" if n >= 8 else "")}
            if not e["active"]:
                r["_tone"] = "muted"
            rows.append(r)
        self.table.set_rows(rows)
        self.count.setText(f"العدد: {digits(len(rows))}")

    # ------------------------------------------------------------- actions
    def _selected(self):
        r = self.table.selected()
        if r is None:
            error(self, "اختر موظفاً من الجدول أولاً.")
        return r

    def _edit_selected(self):
        r = self._selected()
        if r:
            self.edit(r["id"])

    def _enroll_selected(self):
        r = self._selected()
        if r:
            run_enrollment(self, self.ctx, r["id"])
            self.reload()

    def _timeline_selected(self):
        r = self._selected()
        if r:
            self.ctx.window.admin.show_page("timeline", employee_id=r["id"])

    def _delete_selected(self):
        r = self._selected()
        if not r or not confirm(self, f"حذف الموظف «{r['full_name']}» نهائياً مع بصمات وجهه؟\n"
                                      "(لا يمكن حذف موظف له سجل دوام؛ اجعله غير فعّال بدلاً من ذلك)"):
            return
        try:
            self.db.employees.delete(r["id"])
        except StoreError as exc:
            error(self, str(exc))
            return
        self.ctx.worker.reload_galleries()
        self.reload()
        self.toast("تم حذف الموظف")

    def edit(self, employee_id):
        dlg = EmployeeDialog(self.ctx, employee_id, self)
        dlg.exec()
        if dlg.changed:
            self.ctx.worker.reload_galleries()
        self.reload()

    # -------------------------------------------------------------- Excel
    def save_template(self):
        path, _ = QFileDialog.getSaveFileName(self, "حفظ نموذج الاستيراد", "نموذج_الموظفين.xlsx",
                                              "Excel (*.xlsx)")
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "الموظفون"
        ws.sheet_view.rightToLeft = True
        ws.append(IMPORT_HEADERS)
        ws.append(["أحمد علي حسن", "07701234567", "محاسب", self.db.groups.names().get(
            self.db.employees.default_group(), "")])
        for col in "ABCD":
            ws.column_dimensions[col].width = 26
        for r in range(2, 502):
            ws.cell(row=r, column=2).number_format = "@"  # keep the leading zero of phone numbers
        wb.save(path)
        info(self, "تم حفظ النموذج. املأه ثم استورده.")

    def import_excel(self):
        path, _ = QFileDialog.getOpenFileName(self, "استيراد الموظفين", "", "Excel (*.xlsx)")
        if not path:
            return
        try:
            ws = load_workbook(path, data_only=True).active
        except Exception as exc:
            error(self, f"تعذّر فتح الملف: {exc}")
            return
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            error(self, "الملف فارغ.")
            return
        head = [HEADER_ALIASES.get(h, h) for h in (cell_text(c) for c in rows[0])]
        col = {h: i for i, h in enumerate(head) if h}
        if "الاسم" not in col:
            error(self, "يجب أن يحتوي الصف الأول على عمود «الاسم» على الأقل.")
            return
        groups = {g["name"]: g["id"] for g in self.db.groups.list(include_inactive=True)}
        default_group = self.db.employees.default_group()
        by_name = {}
        for e in self.db.employees.list(active=None):
            by_name.setdefault(e["full_name"], e)
        created = updated = 0
        problems = []
        for n, row in enumerate(rows[1:], start=2):
            def get(h):
                i = col.get(h)
                return cell_text(row[i]) if i is not None and i < len(row) else ""
            name, code = get("الاسم"), get("الرقم الوظيفي")
            if not name:
                if code or any(cell_text(v) for v in row):
                    problems.append(f"الصف {n}: لا يوجد اسم")
                continue
            # an existing employee is matched by employee number (old templates) or by full name
            existing = self.db.employees.by_code(code) if code else by_name.get(name)
            data = dict(full_name=name)
            if code and not existing:
                data["code"] = code
            for h, key in (("رقم الهاتف", "phone"), ("الصفة الوظيفية", "job_title")):
                if h in col:
                    data[key] = get(h)
            gname = get("المجموعة")
            if gname and gname not in groups:
                problems.append(f"الصف {n}: المجموعة «{gname}» غير موجودة؛ وُضع في المجموعة الافتراضية")
            if gname in groups:
                data["group_id"] = groups[gname]
            elif not existing and default_group:
                data["group_id"] = default_group
            try:
                eid = self.db.employees.save(data, existing["id"] if existing else None)
                if existing:
                    updated += 1
                else:
                    created += 1
                    by_name.setdefault(name, self.db.employees.get(eid))
            except StoreError as exc:
                problems.append(f"الصف {n}: {exc}")
        self.refresh()
        msg = f"أُضيف {digits(created)} موظفاً وحُدّث {digits(updated)}."
        if problems:
            msg += "\n\nملاحظات:\n" + "\n".join(problems[:20])
        info(self, msg, "نتيجة الاستيراد")


# ---------------------------------------------------------------------------
class EmployeeDialog(QDialog):
    def __init__(self, ctx, employee_id: int | None, parent=None):
        super().__init__(parent)
        self.ctx = ctx
        self.db = ctx.db
        self.employee_id = employee_id
        self.changed = False
        self.setWindowTitle("بيانات الموظف" if employee_id else "إضافة موظف")
        self.setMinimumSize(720, 560)
        lay = QVBoxLayout(self)
        self.tabs = QTabWidget()
        lay.addWidget(self.tabs, 1)
        self._build_info()
        self._build_face()
        self._build_schedule()
        row = QHBoxLayout()
        self.save_btn = button("حفظ", "check", "primary")
        self.save_btn.clicked.connect(self.save)
        close = button("إغلاق")
        close.clicked.connect(self.reject)
        row.addWidget(self.save_btn)
        row.addWidget(close)
        row.addStretch(1)
        lay.addLayout(row)
        self.load()

    # ------------------------------------------------------------- tabs
    def _build_info(self):
        w = QWidget()
        grid = QHBoxLayout(w)
        form = QFormLayout()
        self.name = QLineEdit()
        self.phone = QLineEdit()
        self.job = QLineEdit()
        self.job.setPlaceholderText("مثل: محاسب، مسؤول الشعبة، موظف خدمة")
        self.group = QComboBox()
        self.group_from = DateEdit(self.db.now().date())
        self.active = QCheckBox("فعّال (يمكنه التسجيل في جهاز الحضور)")
        self.pin = QLineEdit()
        self.pin.setEchoMode(QLineEdit.Password)
        self.pin.setPlaceholderText("اتركه فارغاً لعدم التغيير")
        self.code_hint = label("", "Hint")
        form.addRow("الاسم الكامل:", self.name)
        form.addRow("رقم الهاتف:", self.phone)
        form.addRow("الصفة الوظيفية:", self.job)
        # The group is asked only when there is a choice; otherwise everyone joins the default group.
        self.show_group = len(self.db.groups.list(include_inactive=True)) > 1
        g = QHBoxLayout()
        g.addWidget(self.group, 1)
        self.group_from_lbl = label("اعتباراً من")
        g.addWidget(self.group_from_lbl)
        g.addWidget(self.group_from)
        if self.show_group:
            form.addRow("المجموعة:", g)
        else:
            self.group.hide()
            self.group_from_lbl.hide()
            self.group_from.hide()
        if self.employee_id:
            form.addRow("", self.active)
        if self.ctx.setting("kiosk.allow_pin"):
            form.addRow("رمز سري للتسجيل:", self.pin)
            form.addRow("", self.code_hint)
        grid.addLayout(form, 1)
        side = QVBoxLayout()
        self.photo = QLabel(alignment=Qt.AlignCenter)
        self.photo.setFixedSize(150, 170)
        self.photo.setStyleSheet(f"background: {C['muted_bg']}; border-radius: 12px;")
        side.addWidget(self.photo)
        side.addStretch(1)
        grid.addLayout(side)
        self.tabs.addTab(w, "البيانات")

    def _build_face(self):
        w = QWidget()
        lay = QVBoxLayout(w)
        self.face_status = label("", wrap=True)
        lay.addWidget(self.face_status)
        self.samples = QListWidget()
        self.samples.setViewMode(QListWidget.IconMode)
        self.samples.setIconSize(QSize(96, 96))
        self.samples.setResizeMode(QListWidget.Adjust)
        self.samples.setSpacing(6)
        self.samples.setMinimumHeight(260)
        lay.addWidget(self.samples, 1)
        row = QHBoxLayout()
        self.enroll_btn = button("تسجيل الوجه بالكاميرا", "camera", "primary")
        self.enroll_btn.clicked.connect(self._enroll)
        self.photo_btn = button("إضافة من صورة…", "image")
        self.photo_btn.clicked.connect(self._from_photo)
        self.del_sample = button("حذف العينة المحددة", "trash", "danger")
        self.del_sample.clicked.connect(self._delete_sample)
        self.del_all = button("حذف كل العينات", "trash", "danger")
        self.del_all.clicked.connect(self._delete_all)
        for b in (self.enroll_btn, self.photo_btn, self.del_sample, self.del_all):
            row.addWidget(b)
        row.addStretch(1)
        lay.addLayout(row)
        lay.addWidget(label("نصيحة: سجّل الوجه في مكان جهاز الحضور نفسه وبإضاءته المعتادة. يُفضّل ٨ عينات"
                            " فأكثر، وأضف عينات بالنظارة وبدونها إن كان الموظف يستخدمها أحياناً.",
                            "Hint", wrap=True))
        self.tabs.addTab(w, "بصمة الوجه")

    def _build_schedule(self):
        w = QWidget()
        lay = QVBoxLayout(w)
        lay.addWidget(label("جدول دوام خاص بهذا الموظف (يتقدّم على جدول مجموعته خلال الفترة المحددة):",
                            wrap=True))
        self.assign_table = DataTable([("schedule_name", "الجدول", "text"), ("start", "من", "date"),
                                       ("end", "إلى", "date"), ("note", "ملاحظة", "text")],
                                      stretch="schedule_name")
        lay.addWidget(self.assign_table, 1)
        row = QHBoxLayout()
        self.sched_combo = QComboBox()
        self.sched_from = DateEdit(self.db.now().date())
        self.sched_has_to = QCheckBox("حتى")
        self.sched_to = DateEdit(self.db.now().date())
        self.sched_to.setEnabled(False)
        self.sched_has_to.toggled.connect(self.sched_to.setEnabled)
        add = button("إسناد", "plus")
        add.clicked.connect(self._add_assignment)
        rm = button("حذف المحدد", "trash", "danger")
        rm.clicked.connect(self._remove_assignment)
        for x in (label("الجدول:"), self.sched_combo, label("من"), self.sched_from, self.sched_has_to,
                  self.sched_to, add, rm):
            row.addWidget(x)
        row.addStretch(1)
        lay.addLayout(row)
        lay.addWidget(label("سجل المجموعات:", "CardTitle"))
        self.group_hist = DataTable([("name", "المجموعة", "text"), ("start", "من", "date"),
                                     ("end", "إلى", "date")], stretch="name")
        self.group_hist.setMaximumHeight(150)
        lay.addWidget(self.group_hist)
        self.tabs.addTab(w, "الجدول الخاص")

    # ------------------------------------------------------------- data
    def load(self):
        self.group.clear()
        for g in self.db.groups.list(include_inactive=True):
            self.group.addItem(g["name"], g["id"])
        self.sched_combo.clear()
        for s in self.db.schedules.list():
            self.sched_combo.addItem(s["name"], s["id"])
        e = self.db.employees.get(self.employee_id) if self.employee_id else None
        if e is None:
            self.active.setChecked(True)
            self.group.setCurrentIndex(max(0, self.group.findData(self.db.employees.default_group())))
            self.group_from_lbl.hide()
            self.group_from.hide()
            for i in (1, 2):
                self.tabs.setTabEnabled(i, False)
            self.tabs.setTabToolTip(1, "احفظ بيانات الموظف أولاً")
            return
        self.name.setText(e["full_name"])
        self.phone.setText(e["phone"])
        self.job.setText(e["job_title"])
        self.active.setChecked(bool(e["active"]))
        self.code_hint.setText(f"الرقم الوظيفي للتسجيل بالرمز السري: {e['code']}")
        self._orig_group = e["group_id"]
        self.group.setCurrentIndex(max(0, self.group.findData(e["group_id"])))
        self._load_photo()
        self._load_samples()
        self._load_schedule()

    def _load_photo(self):
        data = self.db.employees.photo(self.employee_id) if self.employee_id else None
        if data:
            img = QImage.fromData(data)
            self.photo.setPixmap(QPixmap.fromImage(img).scaled(150, 170, Qt.KeepAspectRatio,
                                                               Qt.SmoothTransformation))
        else:
            self.photo.setPixmap(icon("users", C["muted"], 64).pixmap(64, 64))

    def _load_samples(self):
        self.samples.clear()
        rows = self.db.faces.samples(self.employee_id)
        for s in rows:
            chip = chip_from_png(s["chip"])
            if chip is None:
                continue
            rgb = cv2.cvtColor(chip, cv2.COLOR_BGR2RGB)
            img = QImage(rgb.data, rgb.shape[1], rgb.shape[0], 3 * rgb.shape[1], QImage.Format_RGB888).copy()
            it = QListWidgetItem(QPixmap.fromImage(img), POSE_LABELS.get(s["pose"], s["pose"]))
            it.setData(Qt.UserRole, s["id"])
            it.setToolTip(f"{s['created_at'][:16]} — جودة {s['quality']:.2f}")
            self.samples.addItem(it)
        n = len(rows)
        if n == 0:
            self.face_status.setText("لم يُسجّل وجه هذا الموظف بعد؛ لن يتمكن من التسجيل في جهاز الحضور.")
            self.face_status.setStyleSheet(f"color: {C['bad']}; font-weight: bold;")
        else:
            self.face_status.setText(f"عدد العينات المسجلة: {digits(n)}" +
                                     ("" if n >= 8 else " — يُنصح بإضافة عينات أخرى لدقة أعلى"))
            self.face_status.setStyleSheet(f"color: {C['ok'] if n >= 8 else C['warn']}; font-weight: bold;")

    def _load_schedule(self):
        rows = [dict(a) | {"start": parse_date(a["start_date"]), "end": parse_date(a["end_date"])}
                for a in self.db.schedules.employee_assignments(self.employee_id)]
        self.assign_table.set_rows(rows)
        hist = [dict(h) | {"start": parse_date(h["start_date"]), "end": parse_date(h["end_date"])}
                for h in self.db.employees.group_history(self.employee_id)]
        self.group_hist.set_rows(hist)

    def save(self):
        data = dict(full_name=self.name.text(), phone=self.phone.text(), job_title=self.job.text())
        if self.employee_id:
            data["active"] = self.active.isChecked()
        group = self.group.currentData() if self.show_group else (
            None if self.employee_id else self.db.employees.default_group())
        if group is not None:
            data["group_id"] = group
        effective = self.group_from.value() if self.employee_id else None
        try:
            new_id = self.db.employees.save(data, self.employee_id, group_effective=effective)
            if self.pin.text():
                if not self.pin.text().isdigit() or len(self.pin.text()) < 4:
                    raise StoreError("الرمز السري يجب أن يكون ٤ أرقام على الأقل.")
                self.db.employees.set_pin(new_id, self.pin.text())
        except StoreError as exc:
            error(self, str(exc))
            return
        self.changed = True
        first_time = self.employee_id is None
        self.employee_id = new_id
        if first_time:
            self.setWindowTitle("بيانات الموظف")
            for i in (1, 2):
                self.tabs.setTabEnabled(i, True)
            self.load()
            if confirm(self, "تم حفظ الموظف. هل تريد تسجيل بصمة وجهه الآن؟"):
                self.tabs.setCurrentIndex(1)
                self._enroll()
        else:
            self.ctx.toast("تم الحفظ", "ok")
            self.accept()

    # ------------------------------------------------------------- face
    def _enroll(self):
        if run_enrollment(self, self.ctx, self.employee_id):
            self.changed = True
            self._load_samples()
            self._load_photo()

    def _from_photo(self):
        worker = self.ctx.worker
        if not worker.ready_ok:
            error(self, "نظام التعرّف لم يكتمل تحميله بعد.")
            return
        paths, _ = QFileDialog.getOpenFileNames(self, "اختر صورة أو أكثر للموظف", "",
                                                "الصور (*.jpg *.jpeg *.png *.bmp)")
        if not paths:
            return
        previous = worker.pause()
        try:
            from ...face.enroll import EnrollmentSession

            sess = EnrollmentSession(worker.engine)
            # seed identity check with an existing sample when available
            existing = self.db.faces.samples(self.employee_id)
            problems = []
            for p in paths:
                img = cv2.imdecode(np.fromfile(p, np.uint8), cv2.IMREAD_COLOR)
                msg = sess.add_photo(img) if img is not None else "تعذّرت قراءة الصورة."
                if msg:
                    problems.append(f"{p.split('/')[-1]}: {msg}")
            if sess.captures and existing:
                ref = worker.engine.embed_all(chip_from_png(existing[0]["chip"]))[worker.engine.primary.id]
                sim = float(ref @ sess.captures[0].emb[worker.engine.primary.id])
                if sim < 0.35 and not confirm(self, "الوجه في الصورة لا يشبه العينات المسجلة لهذا الموظف."
                                                    " هل تريد إضافتها رغم ذلك؟"):
                    return
            if sess.captures and not check_duplicates(self, self.ctx, self.employee_id, sess.captures):
                return
            if sess.captures:
                save_captures(self.db, self.employee_id, sess.captures, source="import")
                self.changed = True
                worker.reload_galleries()
            self._load_samples()
            self._load_photo()
            if problems:
                error(self, "\n".join(problems))
        finally:
            worker.set_mode(previous)

    def _delete_sample(self):
        it = self.samples.currentItem()
        if it is None:
            error(self, "اختر عينة أولاً.")
            return
        self.db.faces.delete_sample(it.data(Qt.UserRole))
        self.changed = True
        self._load_samples()
        self.ctx.worker.reload_galleries()

    def _delete_all(self):
        if not confirm(self, "حذف جميع عينات الوجه لهذا الموظف؟ لن يتمكن من التسجيل حتى يُعاد تسجيل وجهه."):
            return
        self.db.faces.delete_all(self.employee_id)
        self.db.employees.set_photo(self.employee_id, None)
        self.changed = True
        self._load_samples()
        self._load_photo()
        self.ctx.worker.reload_galleries()

    # --------------------------------------------------------- schedule
    def _add_assignment(self):
        try:
            self.db.schedules.assign(self.sched_combo.currentData(), self.sched_from.value(),
                                     self.sched_to.value() if self.sched_has_to.isChecked() else None,
                                     employee_id=self.employee_id)
        except StoreError as exc:
            error(self, str(exc))
            return
        self._load_schedule()

    def _remove_assignment(self):
        r = self.assign_table.selected()
        if r and confirm(self, "حذف هذا الإسناد؟"):
            self.db.schedules.delete_assignment(r["id"])
            self._load_schedule()


# ---------------------------------------------------------------------------
def check_duplicates(parent, ctx, employee_id, captures) -> bool:
    """Warn when the new face already belongs to another employee (double enrollment)."""
    engine = ctx.worker.engine
    thr = SECURITY_LEVELS["high"]["sface" if engine.primary.id == "sface-2021dec" else "arc"] + 0.05
    dups = duplicate_candidates(engine, captures, employee_id, thr)
    if not dups:
        return True
    names = ctx.db.employees.names()
    text = "\n".join(f"• {names.get(e, e)} (تشابه {s:.2f})" for e, s in dups[:5])
    return confirm(parent, "تحذير: هذا الوجه يشبه بشدة موظفاً مسجلاً مسبقاً:\n" + text +
                   "\n\nقد يكون الشخص نفسه مسجلاً باسم آخر. هل تريد المتابعة والحفظ؟")


def run_enrollment(parent, ctx, employee_id: int) -> bool:
    worker = ctx.worker
    if not worker.ready_ok:
        error(parent, "نظام التعرّف لم يكتمل تحميله بعد. انتظر قليلاً ثم أعد المحاولة."
              + (f"\n{worker.load_error}" if worker.load_error else ""))
        return False
    dlg = EnrollDialog(ctx, employee_id, parent)
    dlg.exec()
    return dlg.saved


class EnrollDialog(QDialog):
    def __init__(self, ctx, employee_id: int, parent=None):
        super().__init__(parent)
        self.ctx = ctx
        self.db = ctx.db
        self.employee_id = employee_id
        self.saved = False
        emp = self.db.employees.get(employee_id)
        self.setWindowTitle(f"تسجيل وجه: {emp['full_name']}")
        self.setMinimumSize(980, 640)
        lay = QHBoxLayout(self)
        self.view = CameraView(mirror=bool(ctx.setting("camera.mirror")))
        lay.addWidget(self.view, 3)
        side = QVBoxLayout()
        self.instruction = QLabel()
        self.instruction.setWordWrap(True)
        self.instruction.setStyleSheet(f"font-size: 18pt; font-weight: bold; color: {C['navy']};")
        side.addWidget(self.instruction)
        self.hint = label("", wrap=True)
        self.hint.setStyleSheet("font-size: 12pt;")
        side.addWidget(self.hint)
        self.steps = []
        grid = QGridLayout()
        for i, (pose, text, n) in enumerate(STEPS):
            lb = QLabel(f"{digits(i + 1)}. {text.rstrip(' ←→↑↓')}")
            grid.addWidget(lb, i, 0)
            self.steps.append(lb)
        side.addLayout(grid)
        self.bar = QProgressBar()
        self.bar.setRange(0, sum(n for _, _, n in STEPS))
        side.addWidget(self.bar)
        side.addStretch(1)
        self.skip = button("تخطّي هذه الخطوة", kind="ghost")
        self.skip.clicked.connect(self._skip)
        side.addWidget(self.skip)
        self.save_btn = button("حفظ البصمة", "check", "primary")
        self.save_btn.setEnabled(False)
        self.save_btn.clicked.connect(self._save)
        side.addWidget(self.save_btn)
        cancel = button("إلغاء")
        cancel.clicked.connect(self.reject)
        side.addWidget(cancel)
        lay.addLayout(side, 2)

        self.session = ctx.worker.begin_enrollment()
        ctx.worker.enroll_status.connect(self._on_status)
        self.timer = QTimer(self)
        self.timer.timeout.connect(self._frame)
        self.timer.start(33)
        self._seq = -1
        self._update_steps(0)

    def _frame(self):
        frame, seq = self.ctx.camera.latest()
        if frame is None:
            self.view.placeholder = self.ctx.camera.error or "جارٍ تشغيل الكاميرا…"
            self.view.set_frame(None)
        elif seq != self._seq:
            self._seq = seq
            self.view.set_frame(frame)

    def _update_steps(self, current: int):
        for i, lb in enumerate(self.steps):
            if i < current:
                lb.setStyleSheet(f"color: {C['ok']}; font-weight: bold;")
            elif i == current:
                lb.setStyleSheet(f"color: {C['primary']}; font-weight: bold; font-size: 12pt;")
            else:
                lb.setStyleSheet(f"color: {C['muted']};")

    def _on_status(self, st):
        from ...face.pipeline import Status

        tone = st.tone
        self.view.set_status(Status(state="scanning", message="", tone=tone if tone != "busy" else "idle",
                                    box=st.box, progress=st.captured / max(1, st.needed)))
        self.bar.setValue(st.captured)
        if st.done:
            self.instruction.setText("اكتمل التسجيل ✔")
            self.hint.setText("اضغط «حفظ البصمة».")
            self._update_steps(len(STEPS))
            self.save_btn.setEnabled(True)
            self.skip.setEnabled(False)
            self.ctx.worker.set_mode("idle")
            return
        pose, text, _ = STEPS[st.step]
        self.instruction.setText(text)
        self.hint.setText(st.message if st.message != text else "")
        self.hint.setStyleSheet(f"font-size: 12pt; color: {C['bad'] if tone == 'bad' else C['warn'] if tone == 'warn' else C['muted']};")
        self._update_steps(st.step)
        self.skip.setEnabled(pose != "front")
        fronts = sum(1 for c in self.session.captures if c.pose == "front")
        self.save_btn.setEnabled(fronts >= 4)

    def _skip(self):
        previous = self.ctx.worker.pause()
        self.session.skip_step()
        self.ctx.worker.set_mode(previous if previous != "idle" else "enroll")

    def _save(self):
        self.ctx.worker.pause()
        samples = self.session.final_samples()
        if len(samples) < 3:
            error(self, "العينات غير كافية. أعد المحاولة في إضاءة أفضل.")
            return
        if not check_duplicates(self, self.ctx, self.employee_id, samples):
            return
        save_captures(self.db, self.employee_id, samples, "enroll")
        self.ctx.worker.reload_galleries()
        self.saved = True
        self.ctx.toast(f"تم حفظ {digits(len(samples))} عينة لوجه الموظف", "ok")
        self.accept()

    def done(self, r):
        self.timer.stop()
        try:
            self.ctx.worker.enroll_status.disconnect(self._on_status)
        except (RuntimeError, TypeError):
            pass
        self.ctx.worker.set_mode("idle")
        super().done(r)

